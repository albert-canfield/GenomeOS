# SPDX-License-Identifier: AGPL-3.0-or-later
"""The fly's boundary node (genomeos/attribution/simpler_flynode.py, docs/SIMPLER-GENOMES.md section 9).

The boundary rule's cases, the shared boundary bins, the reading of the boundary file, the
populations, the comparison, the rotated-boundary draw and the registered wording are checked on
synthetic inputs. The population, endpoint and matched control are simpler_fly's and are tested in
tests/test_simpler_fly.py.
"""

from __future__ import annotations

import math
from pathlib import Path

import pytest

from genomeos.attribution import simpler
from genomeos.attribution import simpler_fly as sf
from genomeos.attribution import simpler_flynode as fn
from genomeos.attribution.simpler import Chromosome, Gene

ROOT = Path(__file__).resolve().parents[1]


def g(gid: str, start: int, end: int, strand: str) -> Gene:
    return Gene(gid, gid, "chrT", start, end, strand)


# Genes on one synthetic arm: A's 5' end at 10,000, B's at 20,000, C's at 30,000, D's at 60,000.
GENES = [
    g("A", 10_000, 12_000, "+"),
    g("B", 18_000, 20_000, "-"),
    g("C", 30_000, 34_000, "+"),
    g("D", 60_000, 64_000, "+"),
]
CH = Chromosome(GENES)


def test_the_statistics_are_simpler_flys() -> None:
    assert (fn.SEED, fn.ALPHA, fn.MIN_DISCORDANT, fn.CLUSTER_BP) == (
        sf.SEED,
        sf.ALPHA,
        sf.MIN_DISCORDANT,
        sf.CLUSTER_BP,
    )
    assert not hasattr(fn, "domain_rule") and not hasattr(fn, "concordant")
    assert fn.HALF_BIN * 2 == fn.BIN_BP == 500


def test_no_boundary_is_nearest_tss() -> None:
    assert fn.domain_of([], 5) == (-math.inf, math.inf)
    for pos in (0, 14_000, 26_000, 50_000, 90_000):
        gene, case = fn.boundary_rule(CH, [], pos)
        assert gene == CH.nearest_tss(pos) and case == fn.IN_DOMAIN


def test_a_boundary_between_element_and_nearest_moves_the_choice() -> None:
    # Element at 24,000: nearest 5' end is B (20,000, 4 kb). A boundary at 22,000 puts B across it;
    # the domain [22,000, inf) holds C (30,000), which the rule names.
    gene, case = fn.boundary_rule(CH, [22_000], 24_000)
    assert CH.nearest_tss(24_000).gid == "B"
    assert (gene.gid, case) == ("C", fn.MOVES)
    # With the boundary beyond C, B is in the domain again and the rule is nearest TSS.
    gene, case = fn.boundary_rule(CH, [40_000], 24_000)
    assert (gene.gid, case) == ("B", fn.IN_DOMAIN)


def test_an_empty_domain_falls_back_to_nearest_tss() -> None:
    gene, case = fn.boundary_rule(CH, [22_000, 28_000], 25_000)
    assert (gene.gid, case) == (CH.nearest_tss(25_000).gid, fn.EMPTY)


def test_a_boundary_bin_is_shared_by_both_domains() -> None:
    # B's 5' end (20,000) lies 200 bp inside the bin of a boundary at 20,200: in the domain to its right.
    gene, case = fn.boundary_rule(CH, [20_200], 24_000)
    assert (gene.gid, case) == ("B", fn.IN_DOMAIN)
    # At 300 bp it lies outside the bin, across the boundary, and the rule moves to C.
    gene, case = fn.boundary_rule(CH, [20_300], 24_000)
    assert (gene.gid, case) == ("C", fn.MOVES)
    # And from the left, a 5' end inside the right boundary's bin is in the domain.
    lo, hi = fn.domain_of([20_200], 15_000)
    assert lo == -math.inf and hi == 20_450


def _xlsx(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Manual Boundaries"
    ws.append(["Table S4. Manually called boundaries."])
    ws.append(["Chromosome", "Lmost position", "Rmost position"])
    ws.append([None])
    ws.append(["dm6", None, None, "dm3"])
    ws.append(["Chromosome", "Position", None, "Chromosome", "Position"])
    ws.append(["chr2L", 19000, None, "chr2L", 19000])
    ws.append(["chr2L", 19000, None, "chr2L", 19000])  # a repeat collapses
    ws.append(["chrX", 500, None, "Did not liftover", 0])
    ws.append(["chrU", 500, None, "chrU", 500])  # not a used arm
    ws2 = wb.create_sheet("Computational Boundaries")
    ws2.append(["Chromosome", "Lmost position", "Rmost position", "Boundary Score"])
    ws2.append(["dm6", None, None, None])
    ws2.append(["chr3R", 3170000, 3170499, 41.3])
    wb.create_sheet("Representative Boundaries")
    wb.save(path)


def test_boundaries_are_read_as_bin_start_plus_half(tmp_path: Path) -> None:
    p = tmp_path / "b.xlsx"
    _xlsx(p)
    bounds, summary = fn.read_boundaries(p, "manual")
    assert bounds == {"chr2L": [19_250], "chrX": [750]}
    assert summary["boundary_rows"] == 3 and summary["boundaries"] == 2
    comp, s2 = fn.read_boundaries(p, "computational")
    assert comp == {"chr3R": [3_170_250]} and s2["bin_end_minus_start"] == {"499": 1}
    assert fn.read_boundaries(p, "representative")[0] == {}


def _row(mid: int, kind: str = sf.DISTAL, tissue: str = "neural") -> dict:
    e = simpler.Element("chrT", mid - 100, mid + 100, kind, tissue)
    return simpler.attribute([e], {"chrT": CH})[0]


def test_populations_follow_the_named_gene_and_the_host() -> None:
    rows = fn.attribute(
        [_row(24_000), _row(31_000), _row(11_000), _row(20_100, kind=sf.PROMOTER)],
        {"chrT": CH},
        {"chrT": [22_000]},
    )
    between, in_c, in_a, prom = rows
    assert between["domain_case"] != fn.HOST and not between["boundary_gene_contains_element"]
    assert in_c["boundary_gene"] == "C" and in_c["boundary_gene_contains_element"]
    assert [r["mid"] for r in fn.distal(rows)] == [24_000, 31_000, 11_000]
    assert [r["mid"] for r in fn.between_genes(rows)] == [24_000]
    assert [r["mid"] for r in fn.primary(rows)] == [24_000]
    # A nearest gene that contains the element is recorded for the census.
    assert in_a["nearest_gene_contains_element"] and in_a["nearest_gene"] == "A"
    assert prom["kind"] == sf.PROMOTER


def test_census_counts_and_stop_rule() -> None:
    rows = fn.attribute([_row(24_000), _row(26_000), _row(50_000)], {"chrT": CH}, {"chrT": [22_000]})
    cen = fn.census(
        rows, {"other": fn.attribute(rows, {"chrT": CH}, {})}, {"chrT": [22_000]}, {"chrT": 100_000}
    )
    p = cen["primary_outside_the_boundary_gene"]
    assert p["elements"] == 3
    assert p["differ_from_nearest_tss"] == sum(1 for r in rows if r["boundary_gene"] != r["nearest_gene"])
    assert sum(p["boundary_cases"].values()) == p["elements"]
    assert cen["alternative_lists_primary"]["other"]["differ_from_nearest_tss"] == 0
    assert cen["stop_rule_met"] is True  # fewer than 200 differ


def _cpm(tissue: str, value: float = 50.0) -> tuple[float, ...]:
    return tuple(value if t == tissue else 1.0 for t in sf.TISSUES)


def test_compare_counts_both_rules_on_the_same_elements() -> None:
    rows = [
        {"chrom": "chrT", "mid": 1_000, "tissue": "neural", "x": "G1", "y": "G2"},
        {"chrom": "chrT", "mid": 2_500_000, "tissue": "muscle", "x": "G3", "y": "G3"},
    ]
    expr = {"G1": _cpm("neural"), "G2": _cpm("muscle"), "G3": _cpm("muscle")}
    res = fn.compare(rows, expr, "x", "y")
    assert (res["a_concordant"], res["b_concordant"], res["differ"]) == (2, 1, 1)
    assert res["delta_points"] == 50.0 and res["clusters"] == 2
    assert fn.compare([], expr, "x", "y")["delta_ci95_points"] is None


def test_labels_name_the_other_rule() -> None:
    assert fn.label([0.1, 2.0], "NEAREST_AHEAD") == "BOUNDARY_RULE_AHEAD"
    assert fn.label([-2.0, -0.1], "GENE_STRUCTURE_AHEAD") == "GENE_STRUCTURE_AHEAD"
    assert fn.label([-1.0, 1.0], "NEAREST_AHEAD") == "NO_DIFFERENCE_DETECTED"
    assert fn.label(None, "NEAREST_AHEAD") == "NO_DIFFERENCE_DETECTED"


def test_rotation_keeps_the_count_and_reports_the_draws() -> None:
    rows = fn.attribute([_row(24_000)], {"chrT": CH}, {"chrT": [22_000]})
    expr = {"C": _cpm("neural"), "B": _cpm("muscle")}
    out = fn.rotation_null(
        rows, {"chrT": CH}, {"chrT": [22_000]}, {"chrT": 100_000}, expr, observed=1, draws=50
    )
    assert out["draws"] == 50 and 0 <= out["draws_at_or_above_observed"] <= 50
    assert out["p_one_sided"] == round((1 + out["draws_at_or_above_observed"]) / 51, 5)
    again = fn.rotation_null(
        rows, {"chrT": CH}, {"chrT": [22_000]}, {"chrT": 100_000}, expr, observed=1, draws=50
    )
    assert again == out


def _fake_result(lab: str, lab2: str, few: bool) -> dict:
    cmp = {
        "n": 100,
        "differ": 20 if not few else 300,
        "a_concordant": 12,
        "b_concordant": 9,
        "a_rate_pct": 12.0,
        "b_rate_pct": 9.0,
        "delta_points": 3.0,
        "delta_ci95_points": [1.0, 5.0],
    }
    bg = {**cmp, "differ": 90 if few else 250}
    pop = {
        "differ_from_nearest_tss_where_the_nearest_gene_contains_the_element": 7,
        "mde_points_unclustered_vs_nearest_tss": 2.5,
        "gene_structure_names_a_host": 40,
    }
    return {
        "primary": cmp,
        "primary_vs_gene_structure": cmp,
        "between_genes": bg,
        "all_distal": cmp,
        "all_distal_vs_gene_structure": cmp,
        "control": {"expected_rate": 0.03, "p_one_sided": 0.001},
        "rotation_beside": {"mean_rotated_rate_pct": 10.0, "p_one_sided": 0.2},
        "census": {"primary_outside_the_boundary_gene": pop},
        "comparison": lab,
        "comparison_vs_gene_structure": lab2,
        "falsifier": "NOT_FALSIFIED",
        "boundaries": 3121,
    }


@pytest.mark.parametrize("lab", ["BOUNDARY_RULE_AHEAD", "NEAREST_AHEAD", "NO_DIFFERENCE_DETECTED"])
@pytest.mark.parametrize("lab2", ["BOUNDARY_RULE_AHEAD", "GENE_STRUCTURE_AHEAD", "NO_DIFFERENCE_DETECTED"])
def test_reading_fills_the_registered_wording_only(lab: str, lab2: str) -> None:
    text = fn.reading(_fake_result(lab, lab2, few=True))
    w = fn.REGISTRATION["outcome_wording"]
    assert w[lab] in text and w["gene_structure_" + lab2] in text
    assert "{" not in text and "}" not in text
    assert "fewer than the stop rule's 200" in text
    assert text.endswith(w["tail"])
    assert "not shown to carry more" in text  # rotation p 0.2
    assert "fewer than the stop rule" not in fn.reading(_fake_result(lab, lab2, few=False))


def test_registration_words_the_limits() -> None:
    reg = fn.REGISTRATION
    assert "Accessibility is not function" in reg["outcome_wording"]["tail"]
    assert "a boundary is not a regulatory unit by itself" in reg["outcome_wording"]["tail"]
    assert "200" in reg["stop_rule"] and "falsified" in reg["falsifier"]
    assert reg["boundary_set"]["licence"].startswith("CC BY 4.0")


REGISTRATION = ROOT / "data/results/simpler_flynode_registration.json"
INVENTORY = ROOT / "data/results/simpler_flynode_inventory.json"
LOCAL_HOW = (
    "the Stadler 2017 boundary file lives in the git-ignored data/knowledge/simpler_flynode "
    "(scripts/simpler_flynode_run.py fetch) and the Calderon 2022 tables in data/knowledge/simpler_fly"
)


def test_committed_registration_is_the_module_registration() -> None:
    import json

    import tracked_paths as tp

    tp.must_be_present(REGISTRATION, was="registration not yet written")
    reg = json.loads(REGISTRATION.read_text())
    assert reg["written_before_any_rna_value_was_read"] is True
    assert reg["registration"] == json.loads(json.dumps(fn.REGISTRATION))
    cen = reg["census"]
    p, bg, al = cen["primary_outside_the_boundary_gene"], cen["between_genes"], cen["all_distal"]
    for pop in (p, bg, al):
        assert (
            pop["elements"] == sum(pop["elements_by_tissue"].values()) == sum(pop["boundary_cases"].values())
        )
        assert pop["differ_from_nearest_tss"] == pop["boundary_cases"][fn.MOVES]
    # The numbers the registration's text quotes are the census's.
    assert (p["elements"], bg["elements"], al["elements"]) == (1634, 917, 2518)
    assert p["gene_structure_names_a_host"] == 717 and p["elements"] - bg["elements"] == 717
    assert (bg["differ_from_nearest_tss"], p["differ_from_nearest_tss"]) == (90, 234)
    assert p["differ_from_nearest_tss_where_the_nearest_gene_contains_the_element"] == 115
    assert p["differ_from_nearest_tss_where_the_boundary_gene_contains_the_element"] == 0
    assert p["differ_from_nearest_tss"] >= fn.MIN_DISCORDANT and not cen["stop_rule_met"]
    assert reg["boundary_lists"]["manual"]["boundaries"] == 3121
    assert reg["result_manifest"]["complete"] is True and reg["result_manifest"]["code"]["dirty"] is False


def test_inventory_counts_its_downloads_within_budget() -> None:
    import json

    import tracked_paths as tp

    tp.must_be_present(INVENTORY, was="inventory not yet written")
    inv = json.loads(INVENTORY.read_text())
    assert inv["fetched_bytes"] == sum(f["bytes"] for f in inv["fetched"]) == 386_317
    assert inv["lane_download_bytes"] == inv["fetched_bytes"] + inv["scouting_bytes"] <= inv["budget_bytes"]
    assert [f["licence"] for f in inv["fetched"]] == [fn.BOUNDARY_LICENCE]
    assert [f["url"] for f in inv["fetched"]] == [fn.BOUNDARY_URL]
    assert sum(1 for c in inv["considered"] if c["used"]) == 1


@pytest.mark.needs_local_data(str(fn.STORE), how=LOCAL_HOW)
@pytest.mark.needs_local_data(str(sf.STORE), how=LOCAL_HOW)
def test_census_recomputes_from_the_registered_inputs() -> None:
    import json

    reg = json.loads(REGISTRATION.read_text())
    for entry in reg["registered_inputs"].values():
        assert simpler.sha256(ROOT / entry["path"]) == entry["sha256"]
    chroms, tss, _ = sf.read_annotation(ROOT / sf.GFF3)
    calls, _ = sf.tissue_calls(sf.read_markers(ROOT / sf.SUPP), sf.read_cluster_labels(ROOT / sf.SUPP))
    els, _ = sf.elements(calls, sf.read_peaks(ROOT / sf.PEAKS), chroms, tss)
    rows = simpler.attribute(els, chroms)
    sets = {k: fn.read_boundaries(ROOT / fn.BOUNDARY_FILE, k)[0] for k in fn.SHEETS}
    main = fn.attribute(rows, chroms, sets[fn.PRIMARY_SET])
    alts = {k: fn.attribute(rows, chroms, v) for k, v in sets.items() if k != fn.PRIMARY_SET}
    arm = fn.read_arm_lengths(ROOT / sf.GFF3)
    assert fn.census(main, alts, sets[fn.PRIMARY_SET], arm) == reg["census"]


RESULT = ROOT / "data/results/simpler_flynode_result.json"


def test_result_reading_is_the_registered_wording() -> None:
    import json

    import tracked_paths as tp

    tp.must_be_present(RESULT, REGISTRATION, was="result not yet written")
    res = json.loads(RESULT.read_text())
    reg = json.loads(REGISTRATION.read_text())
    assert res["census"] == reg["census"]
    assert res["registration_commit"]
    assert res["positive_control"]["responsive"] is True
    assert (
        res["outcome"] == res["comparison"] == fn.label(res["primary"]["delta_ci95_points"], "NEAREST_AHEAD")
    )
    assert res["comparison_vs_gene_structure"] == fn.label(
        res["primary_vs_gene_structure"]["delta_ci95_points"], "GENE_STRUCTURE_AHEAD"
    )
    assert fn.reading(res) == res["reading"]
    assert (
        res["primary"]["differ"]
        == reg["census"]["primary_outside_the_boundary_gene"]["differ_from_nearest_tss"]
    )
    assert res["attributed_genes_missing_from_rna_rows"] == 0
    assert res["alphagenome_requests"] == 0
    assert res["result_manifest"]["complete"] is True and res["result_manifest"]["code"]["dirty"] is False
